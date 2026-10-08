from pathlib import Path
import re,json,shutil,hashlib,sys
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from appendix_assembly import assemble_technical_appendix

ROOT=Path(__file__).resolve().parent;B=ROOT.parent
NAME=sys.argv[1]
R=ROOT/('build-'+NAME);R.mkdir(exist_ok=True)
OUT=ROOT.parents[1]/'docs/thesis';OUT.mkdir(parents=True,exist_ok=True)
TITLE='FPGAを用いた二次元ゲーム機の設計と実装' if NAME=='本文' else '二次元ゲーム機の設計を再現するための技術付録'
def load_source(name):
 return assemble_technical_appendix(ROOT) if name=='技術付録' else (ROOT/(name+'.md')).read_text(encoding='utf-8')
text=load_source(NAME)
assert not re.search(r'図@|追加図|FIGTOKEN|REFNUM',text)
for rel in sorted(set(re.findall(r'!\[[^\]]+\]\((figs/[^)]+)\)',text))):
 src=ROOT/rel;dst=OUT/rel
 dst.parent.mkdir(parents=True,exist_ok=True)
 if src.exists():shutil.copy2(src,dst)
 assert dst.exists(),rel

# Parse before creating TOCs so that comments in code are not headings.
tokens=[];lines=text.splitlines();i=0;table_number=0;code_number=0;chapter=''
while i<len(lines):
 line=lines[i].strip()
 if not line or line.startswith('# '):i+=1;continue
 if line.startswith('```'):
  language=line[3:];code=[];i+=1
  while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
  assert i<len(lines),'Unclosed code fence'
  code_number+=1
  section=next((x['text'] for x in reversed(tokens) if x['kind']=='heading'),'')
  title=re.sub(r'^(?:\d+|[A-Z])\.\d+\s+','',section)
  label='実行例' if language=='bash' or section.startswith('14.6 ') else ('表示例' if language=='text' else 'コード例')
  tokens.append(dict(kind='code',text='\n'.join(code),language=language,title=f'{label}{chapter}.{code_number}  {title}',label=label));i+=1;continue
 m=re.match(r'^(#{2,5}) (.+)$',line)
 if m:
  level=len(m[1])-1
  # Markdown emphasis and inline-code delimiters describe appearance; they are
  # not part of a Word heading or a table-of-contents label.
  name=re.sub(r'`([^`]+)`|\*\*([^*]+)\*\*',lambda x:x[1] if x[1] is not None else x[2],m[2])
  cm=re.search(r'第(\d+)章|付録([A-Z])',name)
  if cm:
   chapter=cm[1] or cm[2];table_number=0;code_number=0
  sm=re.match(r'^([0-9A-Z]+\.[0-9]+) ',name)
  pm=re.match(r'^第([IVX]+)部 ',name)
  anchor=('s_'+sm[1].replace('.','_') if sm else
          'ch_'+(cm[1] or cm[2]) if cm else
          'part_'+pm[1] if pm else f'h{len(tokens):04d}')
  tokens.append(dict(kind='heading',text=name,level=level,anchor=anchor))
  i+=1;continue
 m=re.match(r'!\[(.+)\]\((.+)\)',line)
 if m:
  tokens.append(dict(kind='image',text=m[1],path=m[2]));assert (OUT/m[2]).exists(),m[2];i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[:\- ]+',v) for v in row):rows.append(row)
   i+=1
  n=len(rows[0]);assert all(len(row)==n for row in rows)
  table_number+=1
  section=next((x['text'] for x in reversed(tokens) if x['kind']=='heading'),'')
  name=re.sub(r'^(?:\d+|[A-Z])\.\d+\s+','',section)
  tokens.append(dict(kind='table',rows=rows,text=f'表{chapter}.{table_number}  {name}'));continue
 group=[line];i+=1
 while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','![','```')):
  group.append(lines[i]);i+=1
 value=''.join(group)
 ref=re.match(r'^\[([A-Za-z][A-Za-z0-9]*)\] (.+)$',value)
 tokens.append(dict(kind='reference',key=ref[1],text=ref[2]) if ref else dict(kind='para',text=value))

heading_anchors=[t['anchor'] for t in tokens if t['kind']=='heading']
assert len(heading_anchors)==len(set(heading_anchors)), 'Duplicate chapter or section bookmark'
reference=OUT/(NAME+'.docx')
doc=Document(reference)
for child in list(doc._element.body):
 if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(25);sec.bottom_margin=Mm(23);sec.left_margin=Mm(25);sec.right_margin=Mm(25)
sec.header_distance=Mm(10);sec.footer_distance=Mm(11)
def font(style,size,bold=False,east='ＭＳ 明朝'):
 style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False;style.font.color.rgb=RGBColor(0,0,0)
 rf=style.element.get_or_add_rPr().find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');style.element.get_or_add_rPr().insert(0,rf)
 rf.set(qn('w:eastAsia'),east)
for name,size,bold,east in [('Normal',11,False,'ＭＳ 明朝'),('Title',22,True,'ＭＳ ゴシック'),('Subtitle',12,False,'ＭＳ 明朝'),('Heading 1',16,True,'ＭＳ ゴシック'),('Heading 2',12,True,'ＭＳ ゴシック'),('Heading 3',11,True,'ＭＳ ゴシック'),('Heading 4',10.5,True,'ＭＳ ゴシック'),('Caption',9.5,True,'ＭＳ ゴシック')]:font(doc.styles[name],size,bold,east)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
pf=doc.styles['Normal'].paragraph_format
pf.line_spacing=Pt(17.5);pf.space_after=Pt(3);pf.first_line_indent=Pt(11);pf.widow_control=True
for name in ['Heading 1','Heading 2','Heading 3','Heading 4']:
 pf=doc.styles[name].paragraph_format;pf.keep_with_next=True;pf.first_line_indent=Pt(0);pf.space_before=Pt(10);pf.space_after=Pt(6)
for name in ['Caption']:
 pf=doc.styles[name].paragraph_format;pf.first_line_indent=Pt(0);pf.line_spacing=Pt(14);pf.space_after=Pt(8)
footer=sec.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title=TITLE;doc.core_properties.author='';doc.core_properties.subject='Wallyと公式RasterIXを用いたFPGAゲーム機の設計と検証'
for t,style in [('2026年度 卒業研究論文' if NAME=='本文' else '卒業研究論文 技術付録','Subtitle'),(TITLE.replace('ための技術付録','ための\n技術付録'),'Title'),('所属・学籍情報 非公開\n研究者（非公開）','Subtitle'),('指導教員 非公開','Subtitle'),(('2026年10月9日 改訂' if NAME=='本文' else '2026年10月8日 改訂'),'Subtitle')]:
 p=doc.add_paragraph(t,style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(30);p.paragraph_format.space_after=Pt(24)

if NAME=='本文' and tokens and tokens[0]['kind']=='para' and tokens[0]['text'].startswith('公開用資料では'):
 note=tokens.pop(0)['text'].replace('`','')
 p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(20)
 for run in p.runs:run.font.size=Pt(9)

toc=[t for t in tokens if t['kind']=='heading' and t['level']<=3]
pages=json.loads((R/'toc-pages.json').read_text(encoding='utf-8')) if (R/'toc-pages.json').exists() else {}
doc.add_heading('目次',1).paragraph_format.page_break_before=True
for entry in toc:
 p=doc.add_paragraph();pf=p.paragraph_format
 pf.first_line_indent=Pt(0);pf.left_indent=Mm((entry['level']-1)*5);pf.space_after=Pt(3 if entry['level']==1 else 0);pf.line_spacing=Pt(15)
 pf.keep_with_next=entry['level']==1;pf.tab_stops.add_tab_stop(Mm(158),WD_TAB_ALIGNMENT.RIGHT)
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),entry['anchor']);r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 size=OxmlElement('w:sz');size.set(qn('w:val'),'21' if entry['level']==1 else '19');pr.append(size)
 if entry['level']==1:pr.append(OxmlElement('w:b'))
 r.append(pr);tt=OxmlElement('w:t');tt.text=entry['text'];r.append(tt);h.append(r);p._p.append(h)
 p.add_run('\t'+str(pages.get(entry['text'],''))).font.size=Pt(9.5)

bid=1
def bookmark(p,name):
 global bid
 a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bid));a.set(qn('w:name'),name)
 b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bid));p._p.insert(0,a);p._p.append(b);bid+=1
reference_keys={t['key'] for t in tokens if t['kind']=='reference'}
def addlink(p,label,target,internal=False):
 h=OxmlElement('w:hyperlink')
 h.set(qn('w:anchor') if internal else qn('r:id'),target if internal else doc.part.relate_to(target,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'1673C3');pr.append(co);r.append(pr)
 v=OxmlElement('w:t');v.text=label;r.append(v);h.append(r);p._p.append(h)
def resolve_ref(label):
 m=re.search(r'([A-Z]|\d+)\.(\d+)',label)
 if m:
  ident=m[1]+'.'+m[2];dest='技術付録' if m[1].isalpha() or '技術付録' in label else '本文'
  raw=load_source(dest)
  if not re.search(r'^#{3,4} '+re.escape(ident)+r' ',raw,re.M):
   if NAME=='技術付録' and dest=='本文':
    dest='技術付録';raw=load_source(dest)
   if not re.search(r'^#{3,4} '+re.escape(ident)+r' ',raw,re.M):return None
  return dest,'s_'+ident.replace('.','_')
 m=re.search(r'技術付録\s*第(\d+)章',label)
 if m:return '技術付録','ch_'+m[1]
 m=re.search(r'第([IVX]+)部',label)
 if m:return '技術付録','part_'+m[1]
 m=re.search(r'付録([A-Z])|第(\d+)章',label)
 if m:return ('技術付録','ch_'+m[1]) if m[1] else ('本文','ch_'+m[2])
 return None

def crosslink(p,label,dest,anchor):
 if dest==NAME:addlink(p,label,anchor,True);return
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),doc.part.relate_to(dest+'.docx',RT.HYPERLINK,is_external=True));h.set(qn('w:anchor'),anchor)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'1673C3');pr.append(co);r.append(pr);v=OxmlElement('w:t');v.text=label;r.append(v);h.append(r);p._p.append(h)

REF_PATTERN=r'技術付録\s*第\d+章|技術付録\s*\d+\.\d+節|第[IVX]+部|付録[A-Z](?:\.\d+)?|第\d+章|(?<![\w.図表])[A-Z]\.\d+(?![\d.])|(?<![\w.図表])\d+\.\d+(?=節|[〜～・]\d+\.\d+節)'

def inline(p,s):
 pattern=r'(`[^`]+`|\*\*[^*]+\*\*|\[[A-Za-z][A-Za-z0-9]*\]|https?://[^\s]+|'+REF_PATTERN+r')'
 parts=[];end=0
 for match in re.finditer(pattern,s):
  if match.start()>end:parts.append((False,s[end:match.start()]))
  parts.append((True,match.group(0)));end=match.end()
 if end<len(s):parts.append((False,s[end:]))
 for matched,bit in parts:
  if bit.startswith('[') and bit[1:-1] in reference_keys:
   addlink(p,bit,'ref_'+bit[1:-1],True);continue
  if bit.startswith(('https://','http://')):
   addlink(p,bit,bit);continue
  if matched and not bit.startswith(('`','**')):
   ref=resolve_ref(bit)
   if ref:crosslink(p,bit,*ref);continue
  r=p.add_run(bit[1:-1] if bit.startswith('`') else bit[2:-2] if bit.startswith('**') else bit)
  if bit.startswith('**'):r.bold=True
  if bit.startswith('`'):r.font.name='Consolas';r.font.size=Pt(9.5)
 return p

def cap(s,center=False):
 p=doc.add_paragraph(style='Caption')
 if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 m=re.match(r'((?:図|表|コード例|実行例|表示例)[\dA-Z案内]+(?:\.\d+)?)(.*)',s)
 if m:p.add_run(m[1]).font.color.rgb=RGBColor.from_string('1673C3');p.add_run(m[2])
 else:p.add_run(s)
 return p
def table(t):
 rows=t['rows'];n=len(rows[0]);p=cap(t['text']);p.paragraph_format.keep_with_next=True
 tab=doc.add_table(rows=1,cols=n);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
 compact_y9=t['text'].startswith('表Y.9 ')
 widths={2:[54,106],3:[47,58,55],4:[43,39,39,39],5:[56,26,26,26,26]}.get(n,[160/n]*n)
 if t['text'].startswith('表案内') and n==3:widths=[55,34,71]
 if rows[0]==['立ち上がり','データ','有効','受取可','この瞬間に渡る語']:widths=[30,28,24,27,51]
 if rows[0]==['境界','運ぶもの','守るべき条件','実装を読む入口']:widths=[32,32,48,48]
 if rows[0]==['設定項目','今回の値','何と一致させるか']:widths=[61,37,62]
 if rows[0]==['resetn','ClockLocked','DDRCalibComplete','この例でのCPUの扱い']:widths=[22,30,43,65]
 if rows[0]==['段階','PSEL','PENABLE','CmdReady','PREADY','FIFOへ入った回数']:widths=[35,15,24,24,25,37]
 if rows[0] in (['1フレーム当たりの時間','1語ずつ','16語FIFO','512語FIFO'], ['全体描画の追加比較','1語ずつ','16語FIFO','512語FIFO']):widths=[61,33,33,33]
 if rows[0]==['区別する版','E6で固定した識別子']:widths=[43,117]
 if rows[0]==['同梱E6内の場所','調べられる内容']:widths=[70,90]
 if rows[0]==['開始ステップ','描画方式','1語ずつ','16語FIFO','512語FIFO']:widths=[25,39,32,32,32]
 if rows[0]==['通常版の確認項目','接続比較版での結果']:widths=[68,92]
 if rows[0]==['版の呼び方','対応する実験記録','固定したソース']:widths=[29,34,97]
 if rows[0]==['記録','使用した版と構成','評価する項目']:widths=[16,57,87]
 if rows[0]==['機能','受け持つこと','本構成での実装']:widths=[35,60,65]
 for c,w in zip(tab.columns,widths):c.width=Mm(w)
 # Do not chain large tables into a single unsplittable page block.
 small=(len(rows)<=7 and sum(sum(len(c) for c in r) for r in rows)<700) or t['text'].startswith(('表14.2','表Y.4'))
 for ri,row in enumerate(rows):
  cells=tab.rows[0].cells if ri==0 else tab.add_row().cells
  for ci,(cell,value) in enumerate(zip(cells,row)):
   cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(1 if compact_y9 else 3);pf.space_after=Pt(1 if compact_y9 else 3);pf.line_spacing=Pt(12.5 if compact_y9 else 13.5)
   pf.keep_with_next=(ri==0 or (small and ri<len(rows)-1))
   inline(p,value)
   for r in p.runs:
    r.font.size=Pt(9.5);r.bold=(ri==0)
    if ri==0:r.font.color.rgb=RGBColor(255,255,255)
   pr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'40' if side in ['top','bottom'] else '80');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar);edges=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single' if side=='bottom' or (ri==0 and side=='top') else 'nil');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'231F20');edges.append(el)
   pr.append(edges);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1673C3' if ri==0 else ('E5EBF8' if ri%2 else 'FFFFFF'));pr.append(sh)
  trpr=tab.rows[ri]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))

section=''
for ti,t in enumerate(tokens):
 kind=t['kind']
 if kind=='heading':
  section=t['text']
  p=doc.add_heading(t['text'],t['level']);bookmark(p,t['anchor'])
  new_part = t['level']==1 and t['text']!='分野別案内'
  new_chapter = (NAME=='技術付録' and t['level']==2 and
                 re.match(r'^(?:第\d+章|付録[A-Z]) ',t['text']))
  if new_part or new_chapter:p.paragraph_format.page_break_before=True
  # Keep the last explanatory unit of these chapters together instead of
  # stranding its closing paragraph on a nearly empty page.
  if t['text'].startswith(('7.11 ','8.8 ')):p.paragraph_format.page_break_before=True
 elif kind=='para':
  p=inline(doc.add_paragraph(),t['text'])
  if section.startswith('3.6 ') and ti+1<len(tokens) and tokens[ti+1]['kind']!='heading':p.paragraph_format.keep_with_next=True
  if section.startswith('8.9 ') and ti+1<len(tokens) and tokens[ti+1]['kind']=='para':p.paragraph_format.keep_with_next=True
  if t['text'].startswith('['):
   p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=Pt(16.5)
   for r in p.runs:r.font.size=Pt(10.5)
 elif kind=='image':
  p=doc.add_paragraph();pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.line_spacing=1;pf.space_after=Pt(0);pf.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.add_run().add_picture(str(OUT/t['path']),width=Mm(158));p._p.xpath('.//wp:docPr')[0].set('descr',t['text']);c=cap(t['text'],True)
  if section.startswith('3.6 '):c.paragraph_format.keep_with_next=True
 elif kind=='table':table(t)
 elif kind=='code':
  heading=cap(t['title']);heading.paragraph_format.keep_with_next=True;heading.paragraph_format.space_after=Pt(4)
  if t['label']=='コード例':
   b=OxmlElement('w:pBdr');e=OxmlElement('w:top');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'8');e.set(qn('w:space'),'5');b.append(e);heading._p.get_or_add_pPr().append(b)
  p=doc.add_paragraph();pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Mm(2);pf.right_indent=Mm(2);pf.line_spacing=Pt(12.5);pf.space_before=Pt(4);pf.space_after=Pt(10);pf.keep_together=len(t['text'].splitlines())<=20
  shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EEF0FF' if t['label']=='実行例' else 'E5EBF8');p._p.get_or_add_pPr().append(shd)
  if t['label']=='コード例':
   b=OxmlElement('w:pBdr');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:space'),'5');b.append(e);p._p.get_or_add_pPr().append(b)
  r=p.add_run(t['text']);r.font.name='Consolas';r.font.size=Pt(9)
  if t['label']=='実行例':r.font.color.rgb=RGBColor.from_string('1673C3')
 elif kind=='reference':
  p=doc.add_paragraph();pf=p.paragraph_format;pf.first_line_indent=Mm(-34);pf.left_indent=Mm(34);pf.tab_stops.add_tab_stop(Mm(34));pf.line_spacing=Pt(16);pf.space_after=Pt(9);pf.keep_together=True
  bookmark(p,'ref_'+t['key']);p.add_run(t['key']).bold=True;p.add_run('\t');inline(p,t['text'])
  for rr in p.runs:rr.font.size=Pt(10.5)

dest=OUT/(NAME+'.docx');doc.save(dest)
# Always remove personal OOXML parts, including any inherited from a template.
sys.path.insert(0,str(ROOT.parents[1]/'tools/publication'))
from sanitize import sanitize_docx,load_config
config=load_config(ROOT.parents[1]/'.private/redactions.json')
dest.write_bytes(sanitize_docx(dest.read_bytes(),config))

alltoc=[t for t in tokens if t['kind']=='heading']
intro='# '+TITLE+'\n\n'
other='技術付録' if NAME=='本文' else '本文'
intro+=f'[資料全体の入口](README.md) ｜ [{other}を開く]({other}.md)'
if NAME == '技術付録':
    intro+=' ｜ [図と数値例で読む技術詳説](../technical/README.md)'
intro+='\n\n## ジャンプできる目次\n\n'
intro+='\n'.join('  '*(t['level']-1)+f'- [{t["text"]}](#{t["anchor"]})' for t in alltoc)+'\n\n'
linked=[];hi=0;fence=False
def md_inline(line):
 parts=re.split(r'(`[^`]+`|!\[.*?\]\([^\n]+\)|https?://[^\s]+)',line)
 for i,part in enumerate(parts):
  if part.startswith(('`','![','https://','http://')):continue
  def subref(m):
   label=m[0];ref=resolve_ref(label)
   if not ref:return label
   dest,anchor=ref
   target=('#'+anchor) if dest==NAME else dest+'.md#'+anchor
   return '['+label+']('+target+')'
  part=re.sub(REF_PATTERN,subref,part)
  part=re.sub(r'\[([A-Za-z][A-Za-z0-9]*)\](?!\()',lambda m:f'[{m[1]}](#ref_{m[1]})' if m[1] in reference_keys else m[0],part)
  parts[i]=part
 return ''.join(parts)
for line in text.splitlines()[1:]:
 if line.startswith('```'):fence=not fence
 if not fence and re.match(r'^#{2,5} ',line):
  linked.append(f'<a id="{alltoc[hi]["anchor"]}"></a>');hi+=1
 elif not fence:
  m=re.match(r'^\[([A-Za-z][A-Za-z0-9]*)\] ',line)
  if m and m[1] in reference_keys:linked.append(f'<a id="ref_{m[1]}"></a>')
  else:line=md_inline(line)
 linked.append(line)
(OUT/(NAME+'.md')).write_text(intro+'\n'.join(linked)+'\n',encoding='utf-8')
(R/'tokens.json').write_text(json.dumps(tokens,ensure_ascii=False,indent=2),encoding='utf-8')
(R/'toc-entries.json').write_text(json.dumps(toc,ensure_ascii=False,indent=2),encoding='utf-8')
stats={'body_characters':len(text),'figures':sum(t['kind']=='image' for t in tokens),'tables':sum(t['kind']=='table' for t in tokens),'code_examples':sum(t['kind']=='code' for t in tokens),'references':len(reference_keys),'toc_entries':len(toc),'all_headings':len(alltoc)}
(R/'document-stats.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
print(json.dumps(stats));print(dest)
