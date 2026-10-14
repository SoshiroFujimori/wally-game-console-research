#!/usr/bin/env python3
"""Build the public abstract while retaining the template's page and text styles."""
import argparse
from copy import deepcopy
import io
import json
from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools/publication"))
from sanitize import load_config, sanitize_docx

p = argparse.ArgumentParser(description=__doc__)
p.add_argument("--template", type=Path, required=True)
p.add_argument("--config", type=Path, required=True)
p.add_argument("--output", type=Path, required=True)
a = p.parse_args()
config = load_config(a.config)
content = json.loads(Path(__file__).with_name("content.json").read_text(encoding="utf-8"))
doc = Document(io.BytesIO(sanitize_docx(a.template.read_bytes(), config)))


def replace(paragraph, value):
    properties = next((deepcopy(r._r.rPr) for r in paragraph.runs if r._r.rPr is not None), None)
    paragraph.clear()
    run = paragraph.add_run(value)
    if properties is not None:
        run._r.insert(0, properties)


# Retain the title, page geometry, columns, figure style, and section styles.
figure_template = deepcopy(next(p._p for p in doc.paragraphs if p._p.findall(".//" + qn("w:drawing"))))
caption_template = deepcopy(next(p._p for p in doc.paragraphs if p.text.startswith("図1")))
original_images = {b.get(qn("r:embed")) for b in doc.element.findall(".//" + qn("a:blip"))}
for paragraph in list(doc.paragraphs):
    if paragraph.style.name == "Title":
        replace(paragraph, content["title"])
    elif paragraph.style.name == "Subtitle":
        paragraph._p.getparent().remove(paragraph._p)

sections = {**content["sections"], "参考文献": content["references"]}
headings = {p.text: p for p in doc.paragraphs if p.style.name == "Heading 1"}
assert set(sections) == set(headings), "Template section headings must match content.json"
figures = content["figures"]
for figure in figures:
    assert figure["section"] in content["sections"]
    assert 1 <= figure["after_paragraph"] <= len(content["sections"][figure["section"]])

for heading, paragraphs in sections.items():
    first = headings[heading]._p
    nodes = []
    node = first.getnext()
    while node is not None and node.tag != qn("w:sectPr"):
        assert node.tag == qn("w:p"), "Unsupported template block"
        wrapper = Paragraph(node, doc._body)
        if wrapper.style.name == "Heading 1":
            break
        nodes.append(node)
        node = node.getnext()
    seed = deepcopy(next(n for n in nodes if Paragraph(n, doc._body).text
                         and not Paragraph(n, doc._body).text.startswith(("図1", "図2"))))
    for node in nodes:
        node.getparent().remove(node)
    anchor = first
    for index, value in enumerate(paragraphs, 1):
        node = deepcopy(seed)
        anchor.addnext(node)
        paragraph = Paragraph(node, doc._body)
        replace(paragraph, value)
        anchor = node
        for number, figure in enumerate(figures, 1):
            if figure["section"] != heading or figure["after_paragraph"] != index:
                continue
            source = (ROOT / figure["source"]).resolve()
            assert source.is_relative_to(ROOT) and source.suffix == ".png"
            picture_node = deepcopy(figure_template)
            anchor.addnext(picture_node)
            picture = Paragraph(picture_node, doc._body)
            picture.clear()
            picture.paragraph_format.keep_with_next = True
            shape = picture.add_run().add_picture(str(source), width=Mm(83))
            shape._inline.docPr.set("descr", figure["alt"])
            caption_node = deepcopy(caption_template)
            picture_node.addnext(caption_node)
            caption = Paragraph(caption_node, doc._body)
            caption.clear()
            caption.paragraph_format.keep_with_next = False
            label, title = figure["caption"].split("  ", 1)
            run = caption.add_run(label + "  ")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("1673C3")
            run = caption.add_run(title)
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            anchor = caption_node

# Discard obsolete embedded images from the retained template.
used_images = {b.get(qn("r:embed")) for b in doc.element.findall(".//" + qn("a:blip"))}
for relationship in original_images - used_images:
    doc.part.drop_rel(relationship)
stream = io.BytesIO()
doc.save(stream)
public = sanitize_docx(stream.getvalue(), config)
a.output.mkdir(parents=True, exist_ok=True)
(a.output / "アブスト_主要事項.docx").write_bytes(public)

md = ["# " + content["title"]]
for heading, paragraphs in content["sections"].items():
    md.append("## " + heading)
    for index, value in enumerate(paragraphs, 1):
        md.append(value)
        for number, figure in enumerate(figures, 1):
            if figure["section"] == heading and figure["after_paragraph"] == index:
                image = a.output / "figs" / f"figure-{number}.png"
                image.parent.mkdir(exist_ok=True)
                image.write_bytes((ROOT / figure["source"]).read_bytes())
                md.append(f"![{figure['alt']}](figs/{image.name})")
                md.append(figure["caption"])
md.append("## 参考文献")
md.extend(content["references"])
md.append("命令の接続設計は[論文第6章](../thesis/本文.md#ch_6)，実機で確認した条件と結果は[第12章](../thesis/本文.md#ch_12)に示す．詳しい実装と比較実験は[技術付録](../thesis/技術付録.md)から参照できる．")
(a.output / "アブスト_主要事項.md").write_bytes(("\n\n".join(md) + "\n").encode("utf-8"))
print(json.dumps({"sections": len(content["sections"]), "figures": len(figures)}))
