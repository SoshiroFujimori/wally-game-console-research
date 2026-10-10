import hashlib,io,json,sys,unittest,zipfile,subprocess,tempfile,shutil
from pathlib import Path
from unittest.mock import patch
from contextlib import redirect_stdout,redirect_stderr
from lxml import etree as E
from PIL import Image,PngImagePlugin
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tools/publication'))
sys.path.insert(0,str(ROOT/'tools/sdwire3'))
import sanitize
import check
import switch

CONFIG={'replacements':{'Example Private Person':'Contributor','例示 著者':'研究者'},'private_terms':['Example Private Person','例示 著者']}

class PublicationTests(unittest.TestCase):
    def setUp(self):check.ERRORS.clear()
    def test_plain_unicode_escaped_and_url(self):
        import urllib.parse
        raw='例示 著者 / '+json.dumps('例示 著者')+' / '+urllib.parse.quote('例示 著者')
        out=sanitize.sanitize_text(raw,CONFIG)
        self.assertNotIn('例示 著者',out);self.assertNotIn('\\u4f8b',out);self.assertIn('研究者',out)

    def test_split_runs_and_private_document_parts(self):
        doc=Document();p=doc.add_paragraph()
        p.add_run('Prefix Example ');p.add_run('Private ').bold=True;p.add_run('Person suffix')
        doc.core_properties.author='Example Private Person'
        doc.core_properties.last_modified_by='Example Private Person'
        mem=io.BytesIO();doc.save(mem)
        before=io.BytesIO()
        with zipfile.ZipFile(mem) as src,zipfile.ZipFile(before,'w') as dst:
            for n in src.namelist():
                if n!='docProps/thumbnail.jpeg':dst.writestr(n,src.read(n))
            dst.writestr('customXml/private.xml','<private>Example Private Person</private>')
            dst.writestr('word/comments.xml','<comments>Example Private Person</comments>')
            dst.writestr('docProps/thumbnail.jpeg',b'private thumbnail')
        result=sanitize.sanitize_docx(before.getvalue(),CONFIG)
        with zipfile.ZipFile(io.BytesIO(result)) as z:
            self.assertFalse(any(n.startswith('customXml/') or 'comments' in n or 'thumbnail' in n for n in z.namelist()))
            body=E.fromstring(z.read('word/document.xml'))
            self.assertIn('Prefix Contributor suffix',''.join(body.itertext()))
            self.assertNotIn(b'Example Private Person',b''.join(z.read(n) for n in z.namelist()))

    def test_replacements_across_multiple_runs(self):
        nodes=[]
        for value in ['Example ','Private Person then ','Example Private',' Person']:
            n=E.Element('t');n.text=value;nodes.append(n)
        sanitize.replace_runs(nodes,CONFIG)
        self.assertEqual(''.join(n.text or '' for n in nodes),'Contributor then Contributor')

    def test_pptx_slide_notes_and_private_parts(self):
        mem=io.BytesIO()
        body=('<a:p xmlns:a="'+sanitize.A+'"><a:r><a:t>Example </a:t></a:r>'
              '<a:r><a:t>Private Person</a:t></a:r></a:p>')
        with zipfile.ZipFile(mem,'w') as z:
            z.writestr('ppt/slides/slide1.xml',body)
            z.writestr('ppt/notesSlides/notesSlide1.xml',body)
            z.writestr('ppt/comments/comment1.xml','<comment>Example Private Person</comment>')
            z.writestr('ppt/commentAuthors.xml','<author>Example Private Person</author>')
            z.writestr('docProps/core.xml','<properties><creator>Example Private Person</creator></properties>')
            z.writestr('docProps/thumbnail.jpeg',b'private thumbnail')
        check.check_file(mem.getvalue(),'sample.pptx',CONFIG,set(),set())
        self.assertTrue(check.ERRORS)
        clean=sanitize.sanitize_bytes(mem.getvalue(),'sample.pptx',CONFIG)
        with zipfile.ZipFile(io.BytesIO(clean)) as z:
            self.assertFalse(any('comment' in n or 'thumbnail' in n for n in z.namelist()))
            for n in ('ppt/slides/slide1.xml','ppt/notesSlides/notesSlide1.xml'):
                self.assertEqual(''.join(E.fromstring(z.read(n)).itertext()),'Contributor')
        check.ERRORS.clear()
        check.check_file(clean,'sample.pptx',CONFIG,set(),set())
        self.assertFalse(check.ERRORS)

    def test_pptx_split_run_privacy_audit(self):
        mem=io.BytesIO()
        with zipfile.ZipFile(mem,'w') as z:
            z.writestr('ppt/notesSlides/notesSlide1.xml',
                '<a:p xmlns:a="'+sanitize.A+'"><a:r><a:t>Example </a:t></a:r>'
                '<a:r><a:t>Private Person</a:t></a:r></a:p>')
        check.check_file(mem.getvalue(),'sample.pptx',CONFIG,set(),set())
        self.assertTrue(check.ERRORS)

    def test_pptx_embedded_object_rejected(self):
        mem=io.BytesIO()
        with zipfile.ZipFile(mem,'w') as z:
            z.writestr('ppt/embeddings/object.bin',b'private data')
        with self.assertRaises(ValueError):
            sanitize.sanitize_bytes(mem.getvalue(),'sample.pptx',CONFIG)

    def test_nested_archive(self):
        inner=io.BytesIO()
        with zipfile.ZipFile(inner,'w') as z:z.writestr('notes.md','Example Private Person')
        outer=io.BytesIO()
        with zipfile.ZipFile(outer,'w') as z:z.writestr('inner.zip',inner.getvalue())
        clean=sanitize.sanitize_bytes(outer.getvalue(),'sample.zip',CONFIG)
        with zipfile.ZipFile(io.BytesIO(clean)) as a,zipfile.ZipFile(io.BytesIO(a.read('inner.zip'))) as b:
            self.assertEqual(b.read('notes.md'),b'Contributor')

    def test_unsafe_zip_path_rejected(self):
        mem=io.BytesIO()
        with zipfile.ZipFile(mem,'w') as z:z.writestr('../escape.txt','hello')
        with self.assertRaises(ValueError):sanitize.sanitize_bytes(mem.getvalue(),'test.zip',CONFIG)

    def test_unknown_binary_rejected(self):
        with self.assertRaises(ValueError):sanitize.sanitize_bytes(b'data','disk.img',CONFIG)

    def test_png_pixels_preserved_metadata_removed(self):
        im=Image.new('RGB',(5,7),(12,34,56));mem=io.BytesIO();meta=PngImagePlugin.PngInfo()
        meta.add_text('Author','Example Private Person');im.save(mem,format='PNG',pnginfo=meta)
        clean=Image.open(io.BytesIO(sanitize.clean_png(mem.getvalue())))
        self.assertEqual(im.tobytes(),clean.tobytes());self.assertNotIn('Author',clean.info)

    def test_uart_nul_is_explicit(self):
        self.assertEqual(sanitize.sanitize_bytes(b'boot\x00ok','uart.raw',CONFIG),b'boot\\x00ok')

    def test_private_values_not_printed(self):
        check.check_text('Example Private Person','notes.md',CONFIG,set())
        self.assertTrue(check.ERRORS)
        self.assertNotIn('Example Private Person',json.dumps(check.ERRORS))

    def test_encoded_private_value_detected(self):
        check.check_text(json.dumps('例示 著者'),'record.json',CONFIG,set())
        self.assertTrue(check.ERRORS)

    def test_unknown_email_detected(self):
        check.check_text('person'+'@'+'example.invalid','notes.md',{},set())
        self.assertTrue(check.ERRORS)

    def test_unknown_image_blocked(self):
        check.check_binary(b'image','image.png',set());self.assertTrue(check.ERRORS)

    def test_known_image_hash_accepted(self):
        check.check_binary(b'image','image.png',{hashlib.sha256(b'image').hexdigest()});self.assertFalse(check.ERRORS)

    def test_machine_name_detected(self):
        check.check_text('DESKTOP-'+'123ABCD','log.txt',{},set());self.assertTrue(check.ERRORS)

    def test_secret_detected_without_leaking_it(self):
        value='ghp_'+'A'*36
        check.check_text(value,'settings.txt',{},set());self.assertTrue(check.ERRORS)
        self.assertNotIn(value,json.dumps(check.ERRORS))

class SDWireTests(unittest.TestCase):
    def test_dry_run_does_not_touch_hardware(self):
        with patch.object(switch.subprocess,'run') as run,patch.object(switch.subprocess,'check_output') as read,redirect_stdout(io.StringIO()) as output:
            self.assertEqual(switch.main(['host','--serial','EXAMPLE_SERIAL']),0)
        run.assert_not_called();read.assert_not_called()
        self.assertFalse(json.loads(output.getvalue())['executed'])

    def test_quiescence_is_required(self):
        with redirect_stderr(io.StringIO()),self.assertRaises(SystemExit):switch.main(['host','--serial','EXAMPLE_SERIAL','--execute'])

    def test_reader_required_before_leaving_host(self):
        with redirect_stderr(io.StringIO()),self.assertRaises(SystemExit):switch.main(['target','--serial','EXAMPLE_SERIAL','--target-quiesced','--execute'])

    def test_mounted_child_rejected(self):
        state={'blockdevices':[{'path':'/dev/example','type':'disk','tran':'usb','mountpoints':[None], 'children':[{'mountpoints':['/media/card']}]}]}
        with self.assertRaises(ValueError):switch.unmounted_usb_disk('/dev/example',state)

    def test_host_exit_queries_children_before_switching(self):
        state={'blockdevices':[{'path':'/dev/example','type':'disk','tran':'usb','mountpoints':[None], 'children':[{'mountpoints':['/media/card']}]}]}
        with patch.object(switch.subprocess,'check_output',return_value=json.dumps(state)) as listing, patch.object(switch.subprocess,'run') as action:
            with self.assertRaises(ValueError):
                switch.main(['target','--serial','EXAMPLE_SERIAL','--reader-device','/dev/example','--target-quiesced','--execute'])
            self.assertIn('--tree',listing.call_args.args[0])
            action.assert_not_called()

    def test_non_usb_rejected(self):
        state={'blockdevices':[{'path':'/dev/example','type':'disk','tran':'nvme','mountpoints':[None]}]}
        with self.assertRaises(ValueError):switch.unmounted_usb_disk('/dev/example',state)

    def test_unmounted_selected_usb_accepted(self):
        state={'blockdevices':[{'path':'/dev/example','type':'disk','tran':'usb','mountpoints':[None], 'children':[{'mountpoints':[None]}]}]}
        switch.unmounted_usb_disk('/dev/example',state)

class HistoryTests(unittest.TestCase):
    def setUp(self):
        self.directory=tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.repo=Path(self.directory.name)
        shutil.copytree(ROOT/'tools/publication',self.repo/'tools/publication',ignore=shutil.ignore_patterns('__pycache__'))
        (self.repo/'.private').mkdir();(self.repo/'publication').mkdir()
        (self.repo/'.private/redactions.json').write_text(json.dumps(CONFIG),encoding='utf-8')
        (self.repo/'publication/policy.json').write_text(json.dumps({'allowed_attribution_emails':['research@example.invalid']}),encoding='utf-8')
        (self.repo/'.gitignore').write_text('.private/\n',encoding='utf-8')
        self.git('init','-q','-b','main')
        self.git('config','user.name','Publication Test')
        self.git('config','user.email','research@example.invalid')

    def git(self,*args):
        return subprocess.check_output(['git','-C',str(self.repo),*args],stderr=subprocess.PIPE).decode().strip()

    def audit(self,*args,input=None):
        return subprocess.run([sys.executable,str(self.repo/'tools/publication/check.py'),*args],input=input,text=True,capture_output=True)

    def test_staged_content_is_checked_even_when_worktree_is_clean(self):
        note=self.repo/'notes.md';note.write_text('Example Private Person',encoding='utf-8')
        self.git('add','notes.md')
        note.write_text('Contributor',encoding='utf-8')
        result=self.audit('--staged')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('private redaction term remains',result.stdout)
        self.assertNotIn('Example Private Person',result.stdout)

    def test_clean_tip_does_not_hide_private_history(self):
        note=self.repo/'notes.md';note.write_text('Example Private Person',encoding='utf-8')
        self.git('add','notes.md');self.git('commit','-qm','Add sample')
        note.write_text('Contributor',encoding='utf-8')
        self.git('add','notes.md');self.git('commit','-qm','Anonymize tip')
        self.assertEqual(self.audit('--staged').returncode,0)
        head=self.git('rev-parse','HEAD')
        result=self.audit('--pre-push',input=f'refs/heads/main {head} refs/heads/main '+('0'*40)+'\n')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('private redaction term remains',result.stdout)

    def test_private_filename_is_not_disclosed(self):
        note=self.repo/'Example Private Person.md';note.write_text('neutral',encoding='utf-8')
        self.git('add',note.name)
        result=self.audit('--staged')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('[private path]',result.stdout)
        self.assertNotIn('Example Private Person',result.stdout)

if __name__=='__main__':unittest.main()
